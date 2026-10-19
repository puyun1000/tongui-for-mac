"""Integration checks using newly generated fixtures; no external documents."""

import tempfile
import zipfile
from pathlib import Path

from docx import Document
from pdf_oxide import DocumentBuilder, EmbeddedFont, PdfDocument

from doc_converter import convert_docx_file
from epub_converter import convert_epub_file
from pdf_converter import convert_pdf_file
from text_converter import convert_ass_file, convert_lrc_file, convert_srt_file, convert_txt_file


def run():
    with tempfile.TemporaryDirectory(prefix='tonggui-verify-') as temp:
        root = Path(temp)
        result = root / 'result'
        result.mkdir()

        plain = root / 'sample.txt'
        plain.write_text('啟淨', encoding='utf-8')
        converted = convert_txt_file(str(plain), str(result), 't2gov')
        assert converted and Path(converted).read_text(encoding='utf-8') == '啓净'
        print('TXT: OK')

        subtitle = root / 'sample.srt'
        subtitle.write_text('1\n00:00:01,000 --> 00:00:02,000\n啟淨\n', encoding='utf-8')
        converted = convert_srt_file(str(subtitle), str(result), 't2gov')
        content = Path(converted).read_text(encoding='utf-8')
        assert '00:00:01,000 --> 00:00:02,000' in content and '啓净' in content
        print('SRT: OK')

        ass = root / 'sample.ass'
        ass.write_text('[Events]\nFormat: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text\n'
                       'Dialogue: 0,0:00:01.00,0:00:02.00,Default,,0,0,0,,啟淨\n', encoding='utf-8')
        converted = convert_ass_file(str(ass), str(result), 't2gov')
        assert converted and '啓净' in Path(converted).read_text(encoding='utf-8')
        print('ASS: OK')

        lyric = root / 'sample.lrc'
        lyric.write_text('[00:01.00]啟淨\n', encoding='utf-8')
        converted = convert_lrc_file(str(lyric), str(result), 't2gov')
        assert converted and '[00:01.00]啓净' in Path(converted).read_text(encoding='utf-8')
        print('LRC: OK')

        document = Document()
        run = document.add_paragraph().add_run('啟淨')
        run.bold = True
        docx_path = root / 'sample.docx'
        document.save(docx_path)
        converted = convert_docx_file(str(docx_path), str(result), 't2gov')
        paragraph = Document(converted).paragraphs[0]
        assert paragraph.text == '啓净' and paragraph.runs[0].bold
        print('DOCX: OK')

        epub_path = root / 'sample.epub'
        with zipfile.ZipFile(epub_path, 'w') as archive:
            archive.writestr('mimetype', 'application/epub+zip', compress_type=zipfile.ZIP_STORED)
            archive.writestr('META-INF/container.xml',
                '<?xml version="1.0"?><container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">'
                '<rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>'
                '</rootfiles></container>')
            archive.writestr('OEBPS/content.opf',
                '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="2.0">'
                '<manifest><item id="chapter" href="chapter.xhtml" media-type="application/xhtml+xml"/></manifest>'
                '<spine><itemref idref="chapter"/></spine></package>')
            archive.writestr('OEBPS/chapter.xhtml',
                '<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml">'
                '<head><title>测验</title></head><body><p>啟淨</p></body></html>')
        converted = convert_epub_file(str(epub_path), str(result), 't2gov')
        with zipfile.ZipFile(converted) as archive:
            assert archive.namelist()[0] == 'mimetype'
            chapter = archive.read('OEBPS/chapter.xhtml').decode('utf-8')
            assert '啟淨' not in chapter and '啓净' in chapter
        print('EPUB: OK')

        font = Path('/System/Library/Fonts/Hiragino Sans GB.ttc')
        if font.exists():
            pdf_path = root / 'sample.pdf'
            (DocumentBuilder()
                .register_embedded_font('CJK', EmbeddedFont.from_file(str(font)))
                .a4_page().font('CJK', 18).at(72, 720).text('啟淨').done()
                .save(str(pdf_path)))
            converted = convert_pdf_file(str(pdf_path), str(result), 't2gov')
            assert converted and '啓净' in PdfDocument(converted).extract_text(0)
            print('PDF text layer: OK')


if __name__ == '__main__':
    run()
